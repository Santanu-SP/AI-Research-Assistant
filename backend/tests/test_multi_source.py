import asyncio
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

from docling.chunking import HybridChunker
from docling.datamodel.base_models import InputFormat
from docling.document_converter import DocumentConverter
import httpx
from pptx import Presentation
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from app.core.errors import AppError
from app.domain.documents import (
    ContentLevel,
    DocumentStatus,
    DocumentType,
    MetadataProvenance,
    SourceType,
)
from app.domain.ingestion import CanonicalMetadata, ParsedDocument, ParsedNode
from app.models.document import Document, DocumentChunk
from app.schemas.retrieval import RetrievalFilters
from app.services import ingestion as ingestion_service
from app.services import sources as source_service
from app.services.doi import normalize_doi
from app.services.docling_ingestion import (
    DoclingIngestionAdapter,
    MarkdownTableSerializerProvider,
)
from app.services.generation import build_grounding_prompt
from app.services.ingestion import stable_node_id
from app.services.scholarly import merge_scholarly_metadata
from app.services.scholarly_providers import (
    CrossrefClient,
    ProviderError,
    crossref_metadata,
    openalex_metadata,
)
from app.services.storage import LocalDocumentStorage
from app.services.url_fetch import SecureUrlFetcher, validate_public_url
from app.schemas.rag import EvidenceItem
from tests.test_ingestion import WordTokenizer


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def make_pptx() -> bytes:
    output = BytesIO()
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "Fixture Presentation"
    title_slide.placeholders[1].text = "Research evidence overview"
    content = presentation.slides.add_slide(presentation.slide_layouts[5])
    content.shapes.title.text = "Results"
    table = content.shapes.add_table(2, 2, 0, 1_000_000, 5_000_000, 1_500_000).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "98%"
    presentation.save(output)
    return output.getvalue()


HTML_FIXTURE = b"""<!doctype html><html><head><meta charset="utf-8"><title>Web Study</title>
<script>window.evil = true;</script></head><body><h1>Web Study</h1>
<h2>Results</h2><p>Transformer web evidence.</p>
<table><tr><th>Metric</th><th>Value</th></tr><tr><td>Accuracy</td><td>98%</td></tr></table>
<a href="https://example.org/reference">Reference</a></body></html>"""

MARKDOWN_FIXTURE = b"""# Markdown Study

## Results

Transformer markdown evidence.

| Metric | Value |
| --- | --- |
| Accuracy | 98% |

[Reference](https://example.org/reference)
"""


class FormatAdapter:
    def parse(self, path: Path, *, document_id: UUID, ingestion_version: str):
        text = f"Structured evidence from {path.suffix}"
        return ParsedDocument(
            document_id=document_id,
            metadata=CanonicalMetadata(title=f"{path.suffix} fixture"),
            parser_name="test-docling",
            ingestion_version=ingestion_version,
            nodes=[
                ParsedNode(
                    node_id=stable_node_id(document_id, 0, text),
                    text=text,
                    chunk_index=0,
                    section="Results",
                    section_path=["Results"],
                )
            ],
        )


async def public_resolver(host: str, port: int):
    if host in {"localhost", "private.example"}:
        return ["127.0.0.1"]
    try:
        return [str(__import__("ipaddress").ip_address(host))]
    except ValueError:
        return ["93.184.216.34"]


def _project_and_user(client) -> tuple[UUID, UUID]:
    project = client.post("/api/v1/projects", json={"name": "Sources"}).json()
    user = client.get("/api/v1/auth/me").json()
    return UUID(project["id"]), UUID(user["id"])


def test_doi_normalization_accepts_common_forms_and_rejects_arbitrary_urls() -> None:
    expected = "10.1000/xyz(123)"
    assert normalize_doi(expected) == expected
    assert normalize_doi(f"doi:{expected}") == expected
    assert normalize_doi(f"https://doi.org/{expected}") == expected
    assert normalize_doi(f"http://dx.doi.org/{expected}") == expected
    try:
        normalize_doi("https://example.org/10.1000/not-a-doi")
    except AppError as exc:
        assert exc.code == "invalid_doi"
    else:
        raise AssertionError("arbitrary URL was accepted as a DOI")


def test_provider_adapters_normalize_metadata_and_abstracts() -> None:
    crossref = crossref_metadata(
        {
            "DOI": "10.1000/FIXTURE",
            "title": ["Deposited title"],
            "author": [{"given": "Ada", "family": "Lovelace"}],
            "published-print": {"date-parts": [[2024, 3, 2]]},
            "container-title": ["Journal of Fixtures"],
            "publisher": "Fixture Press",
            "abstract": "<jats:p>Crossref <b>abstract</b>.</jats:p>",
            "URL": "https://doi.org/10.1000/fixture",
        }
    )
    openalex = openalex_metadata(
        {
            "id": "https://openalex.org/W123",
            "doi": "https://doi.org/10.1000/fixture",
            "display_name": "Indexed title",
            "authorships": [{"author": {"display_name": "Ada Lovelace"}}],
            "publication_year": 2025,
            "publication_date": "2025-01-04",
            "abstract_inverted_index": {"OpenAlex": [0], "abstract": [1]},
            "primary_location": {
                "landing_page_url": "https://doi.org/10.1000/fixture",
                "source": {"display_name": "Journal of Fixtures"},
            },
            "open_access": {"is_oa": True, "oa_status": "gold"},
        }
    )

    assert crossref.doi == "10.1000/fixture"
    assert crossref.authors == ["Ada Lovelace"]
    assert crossref.abstract == "Crossref abstract."
    assert crossref.publisher == "Fixture Press"
    assert openalex.openalex_id == "W123"
    assert openalex.abstract == "OpenAlex abstract"
    assert openalex.open_access_status == "gold"
    assert openalex.field_provenance["title"] is MetadataProvenance.OPENALEX


def test_metadata_merge_is_deterministic_and_records_conflicts() -> None:
    crossref = crossref_metadata(
        {
            "DOI": "10.1000/fixture",
            "title": ["Publisher title"],
            "published": {"date-parts": [[2023]]},
            "publisher": "Fixture Press",
        }
    )
    openalex = openalex_metadata(
        {
            "id": "https://openalex.org/W123",
            "doi": "https://doi.org/10.1000/fixture",
            "display_name": "Different title",
            "publication_year": 2024,
            "open_access": {"oa_status": "green"},
        }
    )
    merged = merge_scholarly_metadata(crossref, openalex)

    assert merged.title == "Publisher title"
    assert merged.publication_year == 2023
    assert merged.openalex_id == "W123"
    assert merged.open_access_status == "green"
    assert merged.field_provenance["title"] is MetadataProvenance.CROSSREF
    assert {item.field for item in merged.conflicts} >= {"title", "publication_year"}


def test_crossref_client_handles_not_found_rate_limit_and_malformed_response() -> None:
    responses = iter(
        [
            httpx.Response(404),
            httpx.Response(429),
            httpx.Response(200, json={"message": "invalid"}),
        ]
    )

    def handler(request: httpx.Request) -> httpx.Response:
        return next(responses)

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    provider = CrossrefClient(
        client,
        base_url="https://api.crossref.org",
        mailto="research@example.org",
        user_agent="Fixture/1",
        timeout=1,
        max_retries=0,
    )
    assert asyncio.run(provider.lookup_doi("10.1000/missing")) is None
    for expected in ("crossref_rate_limited", "crossref_invalid_response"):
        try:
            asyncio.run(provider.lookup_doi("10.1000/fixture"))
        except ProviderError as exc:
            assert exc.code == expected
        else:
            raise AssertionError(f"provider error {expected} was not surfaced")
    asyncio.run(client.aclose())


def test_url_validation_blocks_schemes_and_non_public_addresses() -> None:
    assert asyncio.run(
        validate_public_url("https://example.org/paper#results", public_resolver)
    ) == "https://example.org/paper"
    for url in (
        "file:///etc/passwd",
        "ftp://example.org/paper",
        "http://localhost/paper",
        "http://127.0.0.1/paper",
        "http://10.1.2.3/paper",
        "http://[::1]/paper",
        "http://169.254.169.254/latest/meta-data",
        "http://[fe80::1]/paper",
    ):
        try:
            asyncio.run(validate_public_url(url, public_resolver))
        except AppError as exc:
            assert exc.code in {"unsupported_url_scheme", "unsafe_source_url"}
        else:
            raise AssertionError(f"unsafe URL accepted: {url}")


def test_secure_fetcher_revalidates_redirects_and_enforces_limits() -> None:
    async def redirect_handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(302, headers={"Location": "http://127.0.0.1/admin"})

    redirect_client = httpx.AsyncClient(transport=httpx.MockTransport(redirect_handler))
    fetcher = SecureUrlFetcher(
        redirect_client,
        timeout=2,
        connect_timeout=1,
        max_bytes=32,
        max_redirects=2,
        resolver=public_resolver,
    )
    try:
        asyncio.run(fetcher.fetch("https://example.org/start"))
    except AppError as exc:
        assert exc.code == "unsafe_source_url"
    else:
        raise AssertionError("redirect to loopback was followed")
    asyncio.run(redirect_client.aclose())

    async def oversized_handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, headers={"Content-Type": "text/html"}, content=b"x" * 33)

    client = httpx.AsyncClient(transport=httpx.MockTransport(oversized_handler))
    fetcher = SecureUrlFetcher(
        client, timeout=2, connect_timeout=1, max_bytes=32, max_redirects=0,
        resolver=public_resolver,
    )
    try:
        asyncio.run(fetcher.fetch("https://example.org/large"))
    except AppError as exc:
        assert exc.code == "source_too_large"
    else:
        raise AssertionError("oversized response was accepted")
    asyncio.run(client.aclose())


def test_secure_fetcher_rejects_content_type_and_reports_timeout() -> None:
    def binary_handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, headers={"Content-Type": "application/octet-stream"}, content=b"bin")

    client = httpx.AsyncClient(transport=httpx.MockTransport(binary_handler))
    fetcher = SecureUrlFetcher(
        client, timeout=2, connect_timeout=1, max_bytes=100, max_redirects=0,
        resolver=public_resolver,
    )
    try:
        asyncio.run(fetcher.fetch("https://example.org/file"))
    except AppError as exc:
        assert exc.code == "unsupported_source_content_type"
    else:
        raise AssertionError("binary response was accepted")
    asyncio.run(client.aclose())

    def timeout_handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ReadTimeout("fixture timeout", request=request)

    timeout_client = httpx.AsyncClient(transport=httpx.MockTransport(timeout_handler))
    timeout_fetcher = SecureUrlFetcher(
        timeout_client, timeout=1, connect_timeout=1, max_bytes=100,
        max_redirects=0, resolver=public_resolver,
    )
    try:
        asyncio.run(timeout_fetcher.fetch("https://example.org/slow"))
    except AppError as exc:
        assert exc.code == "source_fetch_timeout"
    else:
        raise AssertionError("timeout was not surfaced")
    asyncio.run(timeout_client.aclose())


def test_official_docling_pipeline_handles_pptx_html_and_markdown(tmp_path: Path) -> None:
    fixtures = {
        "fixture.pptx": make_pptx(),
        "fixture.html": HTML_FIXTURE,
        "fixture.md": MARKDOWN_FIXTURE,
    }
    tokenizer = WordTokenizer()
    adapter = DoclingIngestionAdapter(
        DocumentConverter(
            allowed_formats=[InputFormat.PPTX, InputFormat.HTML, InputFormat.MD]
        ),
        HybridChunker(
            tokenizer=tokenizer,
            serializer_provider=MarkdownTableSerializerProvider(),
        ),
        tokenizer,
    )
    parsed_by_name = {}
    for name, content in fixtures.items():
        path = tmp_path / name
        path.write_bytes(content)
        parsed_by_name[name] = adapter.parse(
            path, document_id=uuid4(), ingestion_version="test-v1"
        )

    assert all(parsed.nodes for parsed in parsed_by_name.values())
    assert any("Accuracy" in node.text for node in parsed_by_name["fixture.pptx"].nodes)
    assert any(
        node.metadata.get("location_kind") == "slide"
        for node in parsed_by_name["fixture.pptx"].nodes
    )
    assert any("Transformer web evidence" in node.text for node in parsed_by_name["fixture.html"].nodes)
    assert any("Transformer markdown evidence" in node.text for node in parsed_by_name["fixture.md"].nodes)


def test_new_local_formats_validate_and_persist_in_owned_project(
    client,
    monkeypatch,
    db_session_factory: sessionmaker[Session],
) -> None:
    monkeypatch.setattr(
        ingestion_service, "ingestion_adapter_for", lambda settings: FormatAdapter()
    )
    project_id, _ = _project_and_user(client)
    fixtures = (
        ("fixture.pptx", make_pptx(), PPTX_MIME, "pptx"),
        ("fixture.html", HTML_FIXTURE, "text/html", "html"),
        ("fixture.md", MARKDOWN_FIXTURE, "text/markdown", "markdown"),
    )
    for name, content, mime, expected_type in fixtures:
        response = client.post(
            f"/api/v1/projects/{project_id}/documents",
            files={"file": (name, content, mime)},
        )
        assert response.status_code == 201, response.text
        assert response.json()["type"] == expected_type
        assert response.json()["projectId"] == str(project_id)

    with db_session_factory() as session:
        documents = list(session.scalars(select(Document)).all())
        assert {document.file_type.value for document in documents} == {
            "pptx", "html", "markdown"
        }
        html = next(document for document in documents if document.file_type.value == "html")
        assert html.checksum


def test_doi_source_persists_abstract_provenance_and_conflicts(
    client,
    db_session: Session,
    test_settings,
) -> None:
    project_id, user_id = _project_and_user(client)
    calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        if "crossref" in request.url.host:
            return httpx.Response(
                200,
                json={
                    "message": {
                        "DOI": "10.1000/fixture",
                        "title": ["Crossref fixture"],
                        "author": [{"given": "Ada", "family": "Lovelace"}],
                        "published": {"date-parts": [[2023]]},
                        "publisher": "Fixture Press",
                    }
                },
            )
        return httpx.Response(
            200,
            json={
                "id": "https://openalex.org/W123",
                "doi": "https://doi.org/10.1000/fixture",
                "display_name": "OpenAlex fixture",
                "publication_year": 2024,
                "abstract_inverted_index": {"Measured": [0], "evidence": [1]},
                "open_access": {"oa_status": "green"},
            },
        )

    async_client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    document = asyncio.run(
        source_service.create_doi_source(
            db_session,
            user_id=user_id,
            project_id=project_id,
            doi="doi:10.1000/FIXTURE",
            settings=test_settings,
            client=async_client,
        )
    )
    assert document.source_type is SourceType.DOI
    assert document.content_level is ContentLevel.ABSTRACT
    assert document.doi == "10.1000/fixture"
    assert document.openalex_id == "W123"
    assert document.chunk_count == 1
    assert document.chunks[0].node_metadata["evidence_scope"] == "abstract"
    assert document.canonical_metadata["metadata_conflicts"]
    assert document.metadata_provenance["title"] == "crossref"

    try:
        asyncio.run(
            source_service.create_doi_source(
                db_session,
                user_id=user_id,
                project_id=project_id,
                doi="10.1000/fixture",
                settings=test_settings,
                client=async_client,
            )
        )
    except AppError as exc:
        assert exc.code == "duplicate_source"
    else:
        raise AssertionError("duplicate DOI was resolved again")
    assert calls == 2
    asyncio.run(async_client.aclose())


def test_metadata_only_source_has_no_evidence_chunks(
    client,
    db_session: Session,
    test_settings,
) -> None:
    project_id, user_id = _project_and_user(client)

    def handler(request: httpx.Request) -> httpx.Response:
        if "crossref" in request.url.host:
            return httpx.Response(
                200,
                json={"message": {"DOI": "10.1000/metadata", "title": ["Metadata only"]}},
            )
        return httpx.Response(404)

    async_client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    document = asyncio.run(
        source_service.create_doi_source(
            db_session,
            user_id=user_id,
            project_id=project_id,
            doi="10.1000/metadata",
            settings=test_settings,
            client=async_client,
        )
    )
    asyncio.run(async_client.aclose())
    assert document.content_level is ContentLevel.METADATA_ONLY
    assert document.status.value == "indexed"
    assert document.chunks == []


def test_url_source_is_project_owned_sanitized_and_deduplicated(
    client,
    db_session: Session,
    test_settings,
    upload_dir: Path,
) -> None:
    project_id, user_id = _project_and_user(client)

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, headers={"Content-Type": "text/html"}, content=HTML_FIXTURE)

    async_client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    fetcher = SecureUrlFetcher(
        async_client,
        timeout=2,
        connect_timeout=1,
        max_bytes=100_000,
        max_redirects=1,
        resolver=public_resolver,
    )
    storage = LocalDocumentStorage(upload_dir, test_settings.document_max_upload_bytes)
    document = asyncio.run(
        source_service.create_url_source(
            db_session,
            user_id=user_id,
            project_id=project_id,
            url="https://example.org/study",
            storage=storage,
            settings=test_settings,
            client=async_client,
            fetcher=fetcher,
            adapter=FormatAdapter(),
        )
    )
    assert document.source_type is SourceType.WEB_PAGE
    assert document.content_level is ContentLevel.WEB_PAGE
    assert document.source_url == "https://example.org/study"
    stored_html = storage.path_for(document.stored_name).read_text()
    assert "<script" not in stored_html
    assert "window.evil" not in stored_html
    assert "Transformer web evidence" in stored_html

    try:
        asyncio.run(
            source_service.create_url_source(
                db_session,
                user_id=user_id,
                project_id=project_id,
                url="https://example.org/copy",
                storage=storage,
                settings=test_settings,
                client=async_client,
                fetcher=fetcher,
                adapter=FormatAdapter(),
            )
        )
    except AppError as exc:
        assert exc.code == "duplicate_source"
    else:
        raise AssertionError("duplicate remote content was indexed twice")
    asyncio.run(async_client.aclose())
    assert len(list(upload_dir.iterdir())) == 1


def test_project_source_listing_filters_and_ownership(
    client,
    db_session: Session,
) -> None:
    project_id, user_id = _project_and_user(client)
    document = Document(
        user_id=user_id,
        project_id=project_id,
        name="Filtered source",
        stored_name=f"metadata-{uuid4()}.json",
        file_type=DocumentType.METADATA,
        mime_type="application/json",
        size=0,
        source_type=SourceType.DOI,
        content_level=ContentLevel.METADATA_ONLY,
        status=DocumentStatus.INDEXED,
        doi="10.1000/filter",
        publication_year=2024,
    )
    db_session.add(document)
    db_session.commit()

    response = client.get(
        f"/api/v1/projects/{project_id}/sources",
        params={
            "sourceType": "doi",
            "contentLevel": "metadata_only",
            "yearFrom": 2024,
            "yearTo": 2024,
            "doi": "https://doi.org/10.1000/filter",
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 1
    assert response.json()["items"][0]["contentLevel"] == "metadata_only"

    # An unrelated project UUID never falls back to a global source list.
    assert client.get(f"/api/v1/projects/{uuid4()}/sources").status_code == 404


def test_typed_retrieval_filters_and_prompt_preserve_content_level() -> None:
    filters = RetrievalFilters(
        sourceTypes=["doi"],
        contentLevels=["abstract"],
        yearFrom=2020,
        yearTo=2025,
        doi="10.1000/fixture",
    )
    assert filters.source_types == [SourceType.DOI]
    evidence = EvidenceItem(
        source_id="S1",
        document_id=uuid4(),
        chunk_id=uuid4(),
        source_type=SourceType.DOI,
        content_level=ContentLevel.ABSTRACT,
        paper_title="Fixture",
        authors=None,
        doi="10.1000/fixture",
        page=None,
        section="Abstract",
        text="Abstract evidence",
        rerank_score=1.0,
    )
    prompt = build_grounding_prompt("Question?", [evidence])
    assert "content_level=abstract" in prompt
    assert "source_type=doi" in prompt
    assert "metadata_only" in prompt
