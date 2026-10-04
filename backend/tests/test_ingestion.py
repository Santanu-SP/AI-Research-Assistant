import asyncio
from io import BytesIO
import json
import logging
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID, uuid4

from docling.chunking import HybridChunker
from docling.datamodel.base_models import InputFormat
from docling.document_converter import DocumentConverter
from docling_core.transforms.chunker.tokenizer.base import BaseTokenizer
from docx import Document as WordDocument
from fastapi import UploadFile
from fastapi.testclient import TestClient
from llama_index.core.schema import Document as LlamaDocument, TextNode
from sqlalchemy import func, select
from sqlalchemy.orm import Session, sessionmaker
from starlette.datastructures import Headers

from app.core.config import Settings
from app.domain.documents import (
    ContentLevel,
    DocumentStatus,
    MetadataProvenance,
    SourceType,
)
from app.domain.ingestion import CanonicalMetadata, ParsedDocument, ParsedNode
from app.models.document import Document, DocumentChunk
from app.schemas.retrieval import RetrievalCandidate
from app.services import documents as document_service
from app.services import ingestion as ingestion_service
from app.services.docling_ingestion import (
    DoclingIngestionAdapter,
    MarkdownTableSerializerProvider,
    canonicalize_docling_output,
)
from app.services.ingestion import DocumentIngestionError, stable_node_id
from app.services.evidence import EvidenceBuilder
from app.services.reranking import RerankerService
from app.services.retrieval import RankedChunk, fuse
from app.services.storage import LocalDocumentStorage


DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


class WordTokenizer(BaseTokenizer):
    max_tokens: int = 512

    def count_tokens(self, text: str) -> int:
        return len(text.split())

    def get_max_tokens(self) -> int:
        return self.max_tokens

    def get_tokenizer(self):
        return self


class FakeDoclingAdapter:
    def __init__(self) -> None:
        self.calls = 0

    def parse(
        self,
        path: Path,
        *,
        document_id: UUID,
        ingestion_version: str,
    ) -> ParsedDocument:
        self.calls += 1
        assert path.suffix == ".docx"
        text = "| Metric | Value |\n| --- | --- |\n| Accuracy | 98% |"
        return ParsedDocument(
            document_id=document_id,
            metadata=CanonicalMetadata(
                title="Fixture Research",
                values={"docling_schema": "DoclingDocument"},
                provenance={"title": MetadataProvenance.DOCLING},
            ),
            parser_name="docling",
            parser_version="test",
            ingestion_version=ingestion_version,
            nodes=[
                ParsedNode(
                    node_id=stable_node_id(document_id, 0, text),
                    text=text,
                    chunk_index=0,
                    section="Results",
                    section_path=["Fixture Research", "Results"],
                    token_count=10,
                    metadata={"content_kind": "table"},
                )
            ],
        )


class FailingAdapter:
    def parse(self, path: Path, *, document_id: UUID, ingestion_version: str):
        raise DocumentIngestionError("The fixture could not be parsed")


class FakeEmbeddingService:
    dimension = 1024

    def __init__(self) -> None:
        self.document_batches: list[list[str]] = []

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        self.document_batches.append(texts)
        return [[1.0] + [0.0] * 1023 for _ in texts]

    def embed_query(self, query: str) -> list[float]:
        return [1.0] + [0.0] * 1023


def make_docx() -> bytes:
    output = BytesIO()
    document = WordDocument()
    document.add_heading("Fixture Research", level=0)
    document.add_heading("Section One", level=1)
    document.add_paragraph("Transformer attention mechanism evidence.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "98%"
    document.add_heading("Section Two", level=1)
    document.add_paragraph("Quantum evidence in the second section.")
    document.save(output)
    return output.getvalue()


def test_canonical_source_and_content_enums_are_stable() -> None:
    assert SourceType.UPLOADED_FILE.value == "uploaded_file"
    assert SourceType.OPENALEX.value == "openalex"
    assert ContentLevel.USER_DOCUMENT.value == "user_document"
    assert ContentLevel.METADATA_ONLY.value == "metadata_only"
    assert MetadataProvenance.DOCLING.value == "docling"


def test_node_identity_is_repeatable_and_content_sensitive() -> None:
    document_id = uuid4()
    first = stable_node_id(document_id, 0, "same text")

    assert first == stable_node_id(document_id, 0, "same text")
    assert first != stable_node_id(document_id, 1, "same text")
    assert first != stable_node_id(document_id, 0, "changed text")


def test_docling_metadata_normalization_preserves_grounding_and_tables() -> None:
    document_id = uuid4()
    raw = {
        "schema_name": "DoclingDocument",
        "version": "1.0.0",
        "origin": {"mimetype": "application/pdf"},
        "pages": {"1": {}, "2": {}},
        "texts": [
            {"label": "title", "text": "Grounded Study"},
            {"label": "section_header", "text": "Abstract"},
            {"label": "text", "text": "Measured abstract evidence."},
            {"label": "section_header", "text": "Results"},
        ],
    }
    llama_document = LlamaDocument(doc_id=str(document_id), text=json.dumps(raw))
    table_text = "| Metric | Value |\n| --- | --- |\n| Accuracy | 98% |"
    node = TextNode(
        text=table_text,
        metadata={
            "doc_items": [
                {
                    "self_ref": "#/tables/0",
                    "label": "table",
                    "prov": [{"page_no": 2}],
                }
            ],
            "headings": ["Grounded Study", "Results"],
            "captions": ["Evaluation metrics"],
        },
    )

    parsed = canonicalize_docling_output(
        document_id=document_id,
        llama_documents=[llama_document],
        llama_nodes=[node],
        tokenizer=WordTokenizer(),
        ingestion_version="test-v1",
    )

    assert parsed.metadata.title == "Grounded Study"
    assert parsed.metadata.abstract == "Measured abstract evidence."
    assert parsed.metadata.provenance["title"] is MetadataProvenance.DOCLING
    assert parsed.page_count == 2
    assert parsed.nodes[0].page == 2
    assert parsed.nodes[0].page_end == 2
    assert parsed.nodes[0].section == "Results"
    assert parsed.nodes[0].section_path == ["Grounded Study", "Results"]
    assert parsed.nodes[0].metadata["content_kind"] == "table"
    assert parsed.nodes[0].metadata["captions"] == ["Evaluation metrics"]
    assert parsed.nodes[0].text == table_text


def test_official_docling_reader_and_node_parser_process_docx(tmp_path: Path) -> None:
    path = tmp_path / "fixture.docx"
    path.write_bytes(make_docx())
    tokenizer = WordTokenizer()
    adapter = DoclingIngestionAdapter(
        DocumentConverter(allowed_formats=[InputFormat.DOCX]),
        HybridChunker(
            tokenizer=tokenizer,
            serializer_provider=MarkdownTableSerializerProvider(),
        ),
        tokenizer,
    )

    parsed = adapter.parse(
        path,
        document_id=uuid4(),
        ingestion_version="test-v1",
    )

    assert parsed.parser_name == "docling"
    assert parsed.metadata.title == "Fixture Research"
    assert any(node.section == "Section One" for node in parsed.nodes)
    table = next(
        node for node in parsed.nodes if node.metadata["content_kind"] == "table"
    )
    assert "Metric" in table.text
    assert "Accuracy" in table.text
    assert table.page is None


def test_legacy_ingestion_fallback_is_explicitly_logged(
    test_settings: Settings,
    caplog,
) -> None:
    settings = test_settings.model_copy(
        update={"document_ingestion_backend": "legacy"}
    )

    with caplog.at_level(logging.WARNING, logger="app.services.ingestion"):
        adapter = ingestion_service.ingestion_adapter_for(settings)

    assert adapter.__class__.__name__ == "LegacyPdfIngestionAdapter"
    assert "legacy pypdf ingestion fallback" in caplog.text


def test_docx_upload_persists_canonical_nodes_and_duplicate_policy(
    client: TestClient,
    monkeypatch,
    upload_dir: Path,
    db_session_factory: sessionmaker[Session],
) -> None:
    adapter = FakeDoclingAdapter()
    monkeypatch.setattr(
        ingestion_service,
        "ingestion_adapter_for",
        lambda settings: adapter,
    )
    content = make_docx()
    alpha = client.post("/api/v1/projects", json={"name": "Alpha"}).json()
    beta = client.post("/api/v1/projects", json={"name": "Beta"}).json()

    first = client.post(
        f"/api/v1/projects/{alpha['id']}/documents",
        files={"file": ("fixture.docx", content, DOCX_MIME)},
    )
    duplicate = client.post(
        f"/api/v1/projects/{alpha['id']}/documents",
        files={"file": ("copy.docx", content, DOCX_MIME)},
    )
    other_project = client.post(
        f"/api/v1/projects/{beta['id']}/documents",
        files={"file": ("fixture.docx", content, DOCX_MIME)},
    )

    assert first.status_code == 201, first.text
    assert first.json()["type"] == "docx"
    assert first.json()["parserName"] == "docling"
    assert first.json()["sourceType"] == "uploaded_file"
    assert first.json()["contentLevel"] == "user_document"
    assert duplicate.status_code == 409
    assert duplicate.json()["error"]["code"] == "duplicate_document"
    assert other_project.status_code == 201
    assert adapter.calls == 2
    assert len(list(upload_dir.iterdir())) == 2

    with db_session_factory() as session:
        documents = list(session.scalars(select(Document)).all())
        chunks = list(session.scalars(select(DocumentChunk)).all())
        assert len(documents) == 2
        assert len(chunks) == 2
        assert all(document.checksum for document in documents)
        assert chunks[0].page is None
        assert chunks[0].section_path == ["Fixture Research", "Results"]
        assert chunks[0].node_metadata["content_kind"] == "table"
        assert chunks[0].node_metadata["title"] == "Fixture Research"


def test_pipeline_batches_existing_embedding_service_and_persists_vector(
    client: TestClient,
    db_session: Session,
    test_settings: Settings,
    upload_dir: Path,
) -> None:
    user_id = UUID(client.get("/api/v1/auth/me").json()["id"])
    settings = test_settings.model_copy(update={"embedding_enabled": True})
    storage = LocalDocumentStorage(upload_dir, settings.document_max_upload_bytes)
    upload = UploadFile(
        file=BytesIO(make_docx()),
        filename="embedded.docx",
        headers=Headers({"content-type": DOCX_MIME}),
    )
    embeddings = FakeEmbeddingService()

    document = asyncio.run(
        document_service.create_document(
            db_session,
            upload,
            storage,
            settings,
            user_id,
            adapter=FakeDoclingAdapter(),
            embeddings=embeddings,
        )
    )

    assert document.status is DocumentStatus.INDEXED
    assert embeddings.document_batches == [[document.chunks[0].text]]
    assert document.chunks[0].embedding is not None
    assert len(document.chunks[0].embedding) == 1024


def test_failed_canonical_ingestion_is_retained_without_partial_chunks(
    client: TestClient,
    monkeypatch,
    db_session_factory: sessionmaker[Session],
) -> None:
    monkeypatch.setattr(
        ingestion_service,
        "ingestion_adapter_for",
        lambda settings: FailingAdapter(),
    )

    response = client.post(
        "/api/v1/documents",
        files={"file": ("failed.docx", make_docx(), DOCX_MIME)},
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "document_processing_failed"
    with db_session_factory() as session:
        document = session.scalar(select(Document))
        assert document is not None
        assert document.status is DocumentStatus.FAILED
        assert document.processing_error == "The fixture could not be parsed"
        assert session.scalar(select(func.count()).select_from(DocumentChunk)) == 0


def test_identical_cross_user_uploads_never_reuse_another_users_document(
    anonymous_client: TestClient,
    monkeypatch,
) -> None:
    adapter = FakeDoclingAdapter()
    monkeypatch.setattr(
        ingestion_service,
        "ingestion_adapter_for",
        lambda settings: adapter,
    )
    content = make_docx()
    created_ids: list[str] = []
    for name in ("Alice", "Bob"):
        email = f"{name.lower()}-duplicate@example.com"
        assert anonymous_client.post(
            "/api/v1/auth/register",
            json={
                "name": name,
                "email": email,
                "password": "long-password-123",
            },
        ).status_code == 201
        assert anonymous_client.post(
            "/api/v1/auth/login",
            json={"email": email, "password": "long-password-123"},
        ).status_code == 200
        project = anonymous_client.post(
            "/api/v1/projects",
            json={"name": f"{name} project"},
        ).json()
        response = anonymous_client.post(
            f"/api/v1/projects/{project['id']}/documents",
            files={"file": ("fixture.docx", content, DOCX_MIME)},
        )
        assert response.status_code == 201
        created_ids.append(response.json()["id"])
        assert anonymous_client.post("/api/v1/auth/logout").status_code == 204

    assert created_ids[0] != created_ids[1]
    assert adapter.calls == 2


def test_newly_ingested_node_remains_compatible_with_rrf_reranking_and_evidence(
    client: TestClient,
    monkeypatch,
    db_session_factory: sessionmaker[Session],
) -> None:
    monkeypatch.setattr(
        ingestion_service,
        "ingestion_adapter_for",
        lambda settings: FakeDoclingAdapter(),
    )
    response = client.post(
        "/api/v1/documents",
        files={"file": ("retrieval.docx", make_docx(), DOCX_MIME)},
    )
    assert response.status_code == 201

    with db_session_factory() as session:
        chunk = session.scalar(select(DocumentChunk))
        assert chunk is not None
        fused = fuse(
            [RankedChunk(chunk=chunk, vector_score=0.9)],
            [RankedChunk(chunk=chunk, keyword_score=0.8)],
            10,
        )
        assert [item.chunk.id for item in fused] == [chunk.id]
        document = chunk.document
        candidate = RetrievalCandidate(
            chunk_id=chunk.id,
            node_id=chunk.node_id,
            document_id=document.id,
            project_id=document.project_id,
            source_type=document.source_type,
            content_level=document.content_level,
            paper_title=document.title,
            authors=document.authors,
            doi=document.doi,
            page=chunk.page,
            page_end=chunk.page_end,
            section=chunk.section,
            section_path=chunk.section_path,
            text=chunk.text,
            vector_score=0.9,
            keyword_score=0.8,
            hybrid_score=2 / 61,
        )

    reranker = RerankerService("fake", batch_size=8)
    reranker._model = SimpleNamespace(predict=lambda pairs, **kwargs: [4.5])
    ranked = reranker.rerank("Which metric?", [candidate])
    evidence = EvidenceBuilder().build(ranked, limit=8)

    assert evidence[0].node_id == candidate.node_id
    assert evidence[0].source_type is SourceType.UPLOADED_FILE
    assert evidence[0].content_level is ContentLevel.USER_DOCUMENT
    assert evidence[0].section_path == ["Fixture Research", "Results"]
